"""テスト. API は呼ばず、リポジトリにコミットしたキャッシュだけを使う."""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from minutes_ai import export
from minutes_ai.evaluate import evaluate, score_meeting
from minutes_ai.extract import Extractor, Usage
from minutes_ai.grounding import check, is_grounded
from minutes_ai.models import ActionItem, Item, Minutes

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / "llm-cache"


def offline(model: str = "claude-haiku-5-5") -> Extractor:
    return Extractor(model=model, cache_dir=CACHE, offline=True)


def test_grounding_ignores_spacing_brackets_and_width():
    source = "佐藤：リリース日は　１１月１６日（月）で確定とします。"
    assert is_grounded("「リリース日は 11月16日（月）で確定」", source)
    assert not is_grounded("リリース日は11月20日で確定", source)
    assert not is_grounded("", source)


def test_check_marks_fabricated_quote():
    minutes = Minutes(
        summary="",
        decisions=[Item("確定", "11月16日で確定"), Item("捏造", "12月1日に延期")],
    )
    check(minutes, "11月16日で確定とします。")
    assert [d.grounded for d in minutes.decisions] == [True, False]


def test_score_matches_owner_alias_and_due():
    gold = {
        "decisions": [],
        "open_issues": [],
        "action_items": [
            {"owner": ["이지은", "イ・ジウン"], "due": "2026-10-16", "keywords": ["仕様書"]},
            {"owner": ["渡辺"], "due": "2026-10-14", "keywords": ["ポート"]},
        ],
    }
    minutes = Minutes(
        summary="",
        action_items=[
            ActionItem("イ・ジウン（韓国開発）", "API仕様書の日本語版を共有", "2026-10-16", "q"),
            ActionItem("渡辺", "ポート開放を申請", "2026-10-15", "q"),
            ActionItem("田中", "社内確認", None, "q"),
        ],
    )
    result = score_meeting("x", minutes, gold)
    assert (result.action_items.matched, result.action_items.predicted) == (2, 3)
    assert result.due_correct == 1
    assert any("期限の誤り" in m for m in result.misses)
    assert any("余分な抽出: 田中" in m for m in result.misses)


def test_offline_mode_raises_without_cache(tmp_path):
    with pytest.raises(LookupError):
        Extractor(cache_dir=tmp_path, offline=True).extract("text", "2026-10-01")


def test_cached_evaluation_replays_without_api():
    extractor = offline()
    results = evaluate(ROOT / "data" / "test", extractor)
    assert len(results) == 6
    assert extractor.usage.calls == 0 and extractor.usage.cache_hits == 6
    assert extractor.usage.estimated_cost_usd > 0


def test_relative_due_date_resolved_from_meeting_date():
    # d05 は火曜の会議. 「明後日」は木曜 (2026-10-22)
    text = (ROOT / "data" / "dev" / "d05_infra_migration.txt").read_text(encoding="utf-8")
    minutes = offline().extract(text, "2026-10-20", "ja")
    yamaguchi = next(a for a in minutes.action_items if "山口" in a.owner)
    assert yamaguchi.due == "2026-10-22"


def test_excel_export_highlights_unverified_rows(tmp_path):
    minutes = Minutes(
        summary="s",
        action_items=[
            ActionItem("鈴木", "準備", "2026-10-09", "q1", grounded=True),
            ActionItem("高橋", "修正", None, "q2", grounded=False),
        ],
    )
    path = tmp_path / "out.xlsx"
    export.to_excel(minutes, path)
    ws = load_workbook(path)["タスク"]
    assert [c.value for c in ws[2]] == ["鈴木", "準備", "2026-10-09", "q1", "OK"]
    assert ws["E3"].value == "要確認"
    assert ws["A3"].fill.fgColor.rgb.endswith("F8D7DA")


def test_markdown_lists_unverified_items():
    minutes = Minutes(summary="s", decisions=[Item("d", "存在しない引用", grounded=False)])
    assert "要確認" in export.to_markdown(minutes)


def test_usage_cost_uses_model_prices():
    usage = Usage()
    usage.add("claude-sonnet-5-5", 1_000_000, 100_000)
    usage.add("claude-haiku-5-5", 1_000_000, 0, cached=True)
    assert usage.cost_usd == pytest.approx(2.0 + 1.0)
    assert usage.estimated_cost_usd == pytest.approx(3.0 + 0.1)


def test_gold_files_are_consistent():
    for gold_path in (ROOT / "data").glob("*/*.gold.json"):
        gold = json.loads(gold_path.read_text(encoding="utf-8"))
        assert gold_path.with_name(gold_path.name.replace(".gold.json", ".txt")).exists()
        for a in gold["action_items"]:
            assert a["owner"] and a["keywords"]
