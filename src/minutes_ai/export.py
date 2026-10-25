"""抽出結果を Excel (タスク管理表) と Markdown (議事録) に出力する."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from minutes_ai.models import Minutes

_HEADER_FILL = PatternFill("solid", fgColor="DDE7F2")
_NG_FILL = PatternFill("solid", fgColor="F8D7DA")


def _mark(grounded: bool | None) -> str:
    return "OK" if grounded else "要確認"


def to_excel(minutes: Minutes, path: str | Path) -> None:
    wb = Workbook()
    sheets = [
        (
            "タスク",
            ["担当", "タスク", "期限", "根拠（原文）", "根拠確認"],
            [
                [a.owner, a.task, a.due or "", a.quote, _mark(a.grounded)]
                for a in minutes.action_items
            ],
            [14, 48, 12, 48, 10],
        ),
        (
            "決定事項",
            ["決定事項", "根拠（原文）", "根拠確認"],
            [[d.text, d.quote, _mark(d.grounded)] for d in minutes.decisions],
            [48, 48, 10],
        ),
        (
            "未決事項",
            ["未決事項", "根拠（原文）", "根拠確認"],
            [[o.text, o.quote, _mark(o.grounded)] for o in minutes.open_issues],
            [48, 48, 10],
        ),
    ]
    wb.remove(wb.active)
    for title, header, rows, widths in sheets:
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
        for row in rows:
            ws.append(row)
            # 原文に根拠が見つからない行は色を付けて、人が確認すべき箇所を示す
            if row[-1] != "OK":
                for cell in ws[ws.max_row]:
                    cell.fill = _NG_FILL
        for i, width in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + i)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    wb.save(path)


def to_markdown(minutes: Minutes, title: str = "議事録") -> str:
    lines = [f"# {title}", "", "## 要約", "", minutes.summary, "", "## 決定事項", ""]
    lines += [f"- {d.text}" for d in minutes.decisions] or ["- なし"]
    lines += ["", "## タスク", "", "| 担当 | タスク | 期限 |", "|---|---|---|"]
    lines += [f"| {a.owner} | {a.task} | {a.due or '—'} |" for a in minutes.action_items]
    lines += ["", "## 未決事項", ""]
    lines += [f"- {o.text}" for o in minutes.open_issues] or ["- なし"]
    unverified = [i for i in minutes.all_items() if not i.grounded]
    if unverified:
        lines += ["", "## 要確認（根拠が原文に見つからない項目）", ""]
        lines += [f"- {i.quote}" for i in unverified]
    return "\n".join(lines) + "\n"
